import os
import io
import base64
import pandas as pd
import re
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from openpyxl import load_workbook
from puma_project_resolver import resolve_okd_project, clean_subject_project
from puma_runtime_config import test_safe_env

# ============================ Config ============================
SPREADSHEET_ID = test_safe_env('PUMA_SPREADSHEET_ID', '1pwVlYSGVjyTCLt4GT7xU2TCnxfdJuxAbp_jU6Snisls')
SCOPES = [
    'https://www.googleapis.com/auth/gmail.readonly',
    'https://www.googleapis.com/auth/gmail.modify',
    'https://www.googleapis.com/auth/spreadsheets'
]
DEBUG = True  # set False to quiet debug prints

# ====================== Auth / Services ========================
def setup_services():
    creds = None
    if os.path.exists('token.json'):
        creds = Credentials.from_authorized_user_file('token.json', SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            if not os.path.exists('client_secrets.json'):
                print("Error: 'client_secrets.json' not found.")
                return None, None
            flow = InstalledAppFlow.from_client_secrets_file('client_secrets.json', SCOPES)
            creds = flow.run_local_server(port=0)
        with open('token.json', 'w') as token:
            token.write(creds.to_json())
    try:
        return build('gmail', 'v1', credentials=creds), build('sheets', 'v4', credentials=creds)
    except HttpError as e:
        print(f"Auth error: {e}")
        return None, None

# ===================== Header utilities ========================
def _n(x): return (str(x) if x is not None else "").strip().lower()
NEEDED_KEYS = ("Quantity", "Type", "Manufacturer", "Part Number")

def _find_header_map_on_row(values):
    """Detect our 4 columns on a header row. Supports combined 'Manufacturer / Product Number'."""
    idx_map = {}
    for idx, v in enumerate(values):
        t = _n(v)
        if t in {"qty", "quantity", "qnty", "qty ordered", "qty/ea"}:
            idx_map["Quantity"] = idx
        elif t in {"type", "item type", "category"}:
            idx_map["Type"] = idx
        elif ("manufacturer" in t) and ("product" in t):
            idx_map["Manufacturer"] = idx
            idx_map["Part Number"] = idx + 1
        elif t in {"manufacturer", "mfr", "mfg", "brand", "vendor"}:
            idx_map["Manufacturer"] = idx
        elif t in {"part number", "part #", "part no", "item #", "sku", "item", "item no"}:
            idx_map["Part Number"] = idx
    score = sum(1 for k in NEEDED_KEYS if k in idx_map)
    return idx_map if score >= 3 else None

# ================== XLSX / CSV readers (master sheet) ==========
def _read_master_sheet_xlsx(file_stream, project_name, filename):
    wb = load_workbook(file_stream, data_only=True)
    ws = None
    for s in wb.worksheets:
        if _n(s.title) == "master quotation":
            ws = s
            break
    if ws is None:
        ws = wb.worksheets[0]
        if DEBUG:
            print(f"[DEBUG] '{filename}': 'master quotation' not found; using '{ws.title}'")

    header_row, header_map = None, None
    max_check = min(200, ws.max_row or 200)
    for r in range(1, max_check + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if not any(vals):
            continue
        cand = _find_header_map_on_row(vals)
        if cand:
            header_row, header_map = r, cand
            break

    if not header_map:
        if DEBUG:
            print(f"[DEBUG] '{filename}': no usable header row on '{ws.title}'")
        return []

    if DEBUG:
        print(f"[DEBUG] '{filename}': using '{ws.title}' header row {header_row} map {header_map}")

    rows = []
    r = header_row + 1
    empty_streak = 0
    while r <= (ws.max_row or r + 500):
        row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)] if r <= ws.max_row else []
        if not row or not any(row):
            empty_streak += 1
            if empty_streak >= 3:
                break
            r += 1
            continue
        empty_streak = 0

        def pick(key, default=""):
            idx = header_map.get(key)
            if idx is None or idx >= len(row):
                return default
            v = row[idx]
            return "" if v is None else v

        # Default each imported line to Unapproved
        line = {
            "Project": project_name.strip(),
            "Source": filename,
            "Type": pick("Type", ""),
            "Part Number": str(pick("Part Number", "")).strip(),
            "Manufacturer": pick("Manufacturer", ""),
            "Quantity": pick("Quantity", ""),
            "Status": "Unapproved",
            "PO Number": "",
            "Estimated Ship Date (ESD)":"",
            "Date Received":"",
            "Date Scheduled":"",
            "Date Delivered":""
        }
        if any([line["Part Number"], line["Manufacturer"], line["Type"]]):
            rows.append(line)
        r += 1
    return rows

def _read_master_sheet_csv(file_bytes, project_name, filename):
    df = pd.read_csv(io.BytesIO(file_bytes))
    cols = list(df.columns)
    rename = {}
    for i, c in enumerate(cols):
        t = _n(c)
        if t in {"qty", "quantity", "qnty", "qty ordered", "qty/ea"}:
            rename[c] = "Quantity"
        elif t in {"type", "item type", "category"}:
            rename[c] = "Type"
        elif ("manufacturer" in t) and ("product" in t):
            rename[c] = "Manufacturer"
            if i + 1 < len(cols):
                rename[cols[i + 1]] = "Part Number"
        elif t in {"manufacturer", "mfr", "mfg", "brand", "vendor"}:
            rename[c] = "Manufacturer"
        elif t in {"part number", "part #", "part no", "item #", "sku", "item", "item no"}:
            rename[c] = "Part Number"
    if rename:
        df = df.rename(columns=rename)

    out = []
    for _, rec in df.iterrows():
        line = {
            "Project": project_name.strip(),
            "Source": filename,
            "Type": rec.get("Type", ""),
            "Part Number": str(rec.get("Part Number", "")).strip(),
            "Manufacturer": rec.get("Manufacturer", ""),
            "Quantity": rec.get("Quantity", ""),
            "Status": "Unapproved",
            "PO Number": "",
            "Estimated Ship Date (ESD)":"",
            "Date Received":"",
            "Date Scheduled":"",
            "Date Delivered":""
        }
        if any([line["Part Number"], line["Manufacturer"], line["Type"]]):
            out.append(line)
    return out

# =========================== Gmail =============================
def get_emails(gmail_service):
    print("Fetching OKD emails from PUMA inbox...")
    out = []
    req = gmail_service.users().messages().list(
        userId="me",
        labelIds=[test_safe_env('PUMA_OKD_LABEL_ID', 'Label_1716122347040870890')],
        q="is:unread has:attachment",
        maxResults=100
    )
    while req is not None:
        resp = req.execute()
        out.extend(resp.get("messages", []))
        req = gmail_service.users().messages().list_next(previous_request=req, previous_response=resp)
    print(f"Found {len(out)} unread OKD email(s).") if out else print("No new OKD emails found.")
    return out


OKD_PREFIXES = ("OKD -", "OKD –", "OKD—", "OKD–", "OKD")  # guard variants

def _is_okd_subject(s: str) -> bool:
    if not s:
        return False
    su = s.strip().upper()
    return su.startswith("OKD")  # starts with OKD in any of the above variants
def get_and_process_emails_single(gmail_service, message_ref):
    try:
        msg = gmail_service.users().messages().get(userId="me", id=message_ref["id"], format="full").execute()
        payload = msg.get("payload", {})
        headers = payload.get("headers", [])
        subject = next((h["value"] for h in headers if h["name"] == "Subject"), "")

        # OKD is the one flow allowed to create a new project, but the subject
        # still must resolve against an existing tracker or Open Projects entry.
        project_name = clean_subject_project(subject, ["OKD", "Okay to Design"])

        parsed = []

        def walk(part):
            if "parts" in part:
                for p in part["parts"]:
                    yield from walk(p)
            else:
                yield part

        for part in walk(payload):
            filename = part.get("filename") or ""
            if not filename or not filename.lower().endswith((".xlsx", ".csv")):
                continue
            attach_id = part.get("body", {}).get("attachmentId")
            if not attach_id:
                continue
            att = gmail_service.users().messages().attachments().get(
                userId="me", messageId=message_ref["id"], id=attach_id
            ).execute()
            file_data = base64.urlsafe_b64decode(att["data"])
            file_stream = io.BytesIO(file_data)

            if filename.lower().endswith(".xlsx"):
                rows = _read_master_sheet_xlsx(file_stream, project_name, filename)
            else:
                rows = _read_master_sheet_csv(file_data, project_name, filename)
            parsed.extend(rows)

        print(f"Parsed {len(parsed)} row(s) from: {subject}")
        return project_name, parsed

    except HttpError as e:
        print(f"Gmail API error on {message_ref.get('id')}: {e}")
        return "(unknown project)", []
    except Exception as e:
        print(f"Unexpected error on {message_ref.get('id')}: {e}")
        return "(unknown project)", []

# ========================== Sheets =============================
def ensure_dashboard_exists(sheets_service):
    meta = sheets_service.spreadsheets().get(spreadsheetId=SPREADSHEET_ID).execute()
    sheets = meta.get('sheets', [])
    titles = {s['properties']['title']: s['properties']['sheetId'] for s in sheets}
    if "Dashboard" in titles:
        return titles["Dashboard"]
    resp = sheets_service.spreadsheets().batchUpdate(
        spreadsheetId=SPREADSHEET_ID, body={'requests':[{'addSheet':{'properties':{'title':'Dashboard'}}}]}
    ).execute()
    dash_id = resp['replies'][0]['addSheet']['properties']['sheetId']
    sheets_service.spreadsheets().values().update(
        spreadsheetId=SPREADSHEET_ID, range="Dashboard!A1:H1", valueInputOption='RAW',
        body={'values': [["Project Name","Open Tasks","Total Line Items","Unapproved","Approved","Ordered","Received","Delivered"]]}
    ).execute()
    sheets_service.spreadsheets().batchUpdate(
        spreadsheetId=SPREADSHEET_ID,
        body={'requests':[{'updateSheetProperties':{'properties':{'sheetId':dash_id,'gridProperties':{'frozenRowCount':1}},'fields':'gridProperties.frozenRowCount'}}]}
    ).execute()
    return dash_id

def _find_dashboard_row(sheets_service, project_name, tracker_sheet_id):
    """Return 1-based row index in Dashboard where this project lives, else None."""
    res = sheets_service.spreadsheets().values().get(
        spreadsheetId=SPREADSHEET_ID, range="Dashboard!A:A", valueRenderOption="FORMULA"
    ).execute()
    rows = res.get("values", []) or []
    target_gid = f"#gid={tracker_sheet_id}"
    for i, r in enumerate(rows, start=1):
        val = r[0] if r else ""
        if isinstance(val, str):
            if val.startswith("=HYPERLINK(") and target_gid in val:
                return i
            if val.strip() == project_name.strip():
                return i
    return None

def upsert_dashboard_row(sheets_service, project_name, tracker_sheet_id):
    """Ensure Dashboard has link in A and summary formulas in B..H for this project."""
    ensure_dashboard_exists(sheets_service)
    row = _find_dashboard_row(sheets_service, project_name, tracker_sheet_id)

    if not row:
        res = sheets_service.spreadsheets().values().get(
            spreadsheetId=SPREADSHEET_ID, range="Dashboard!A:A"
        ).execute()
        rows = res.get('values', [])
        row = len(rows) + 1 if rows else 2

    url = f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/edit#gid={tracker_sheet_id}"
    link = f'=HYPERLINK("{url}","{project_name}")'

    rng_tasks  = f"'{project_name} - Tasks'!B2:B"
    rng_itemsB = f"'{project_name} - Project Tracker'!B2:B"
    rng_status = f"'{project_name} - Project Tracker'!G2:G"

    open_tasks   = f"=COUNTA({rng_tasks})"
    total_items  = f"=COUNTA({rng_itemsB})"
    unapproved   = f"=COUNTIF({rng_status},\"Unapproved\")"
    approved     = f"=COUNTIF({rng_status},\"Approved\")"
    ordered      = f"=COUNTIF({rng_status},\"Ordered\")"
    received     = f"=COUNTIF({rng_status},\"Received\")"
    delivered    = f"=COUNTIF({rng_status},\"Delivered\")"

    sheets_service.spreadsheets().values().update(
        spreadsheetId=SPREADSHEET_ID,
        range=f"Dashboard!A{row}:H{row}",
        valueInputOption="USER_ENTERED",
        body={"values": [[link, open_tasks, total_items, unapproved, approved, ordered, received, delivered]]}
    ).execute()

def apply_dropdowns(sheets_service, titles_map, tab_tracker, tab_tasks):
    """Add data validation dropdowns for Status columns in both tabs (idempotent)."""
    tracker_sheet_id = titles_map[tab_tracker]
    tasks_sheet_id = titles_map[tab_tasks]

    tracker_status_rule = {
        'setDataValidation': {
            'range': {'sheetId': tracker_sheet_id, 'startRowIndex': 1, 'endRowIndex': 2000, 'startColumnIndex': 6, 'endColumnIndex': 7},
            'rule': {
                'condition': {'type': 'ONE_OF_LIST',
                              'values': [{'userEnteredValue': v} for v in ["Unapproved","Approved","Ordered","Received","Scheduled","Delivered"]]},
                'showCustomUi': True, 'strict': True
            }
        }
    }

    tasks_status_rule = {
        'setDataValidation': {
            'range': {'sheetId': tasks_sheet_id, 'startRowIndex': 1, 'endRowIndex': 2000, 'startColumnIndex': 4, 'endColumnIndex': 5},
            'rule': {
                'condition': {'type': 'ONE_OF_LIST',
                              'values': [{'userEnteredValue': v} for v in ["Open","Closed"]]},
                'showCustomUi': True, 'strict': True
            }
        }
    }

    sheets_service.spreadsheets().batchUpdate(
        spreadsheetId=SPREADSHEET_ID, body={'requests': [tracker_status_rule, tasks_status_rule]}
    ).execute()

def apply_status_colors(sheets_service, titles_map, tab_tracker, tab_tasks, add_for_tracker, add_for_tasks):
    """Add conditional formatting color rules (only when a sheet is newly created)."""
    requests = []

    def color(r, g, b):
        return {'red': r, 'green': g, 'blue': b}

    if add_for_tracker:
        tracker_sheet_id = titles_map[tab_tracker]
        base_range = {'sheetId': tracker_sheet_id, 'startRowIndex': 1, 'endRowIndex': 2000, 'startColumnIndex': 6, 'endColumnIndex': 7}
        status_colors = {
            "Unapproved": color(1.0, 0.85, 0.85),
            "Approved":   color(0.80, 0.88, 1.00),
            "Ordered":    color(1.00, 0.90, 0.60),
            "Received":   color(0.75, 0.92, 0.92),
            "Scheduled":  color(0.92, 0.85, 1.00),
            "Delivered":  color(0.80, 1.00, 0.80),
        }
        for label, bg in status_colors.items():
            requests.append({
                'addConditionalFormatRule': {
                    'index': 0,
                    'rule': {
                        'ranges': [base_range],
                        'booleanRule': {
                            'condition': {'type': 'TEXT_EQ', 'values': [{'userEnteredValue': label}]},
                            'format': {'backgroundColor': bg, 'textFormat': {'bold': True}}
                        }
                    }
                }
            })

    if add_for_tasks:
        tasks_sheet_id = titles_map[tab_tasks]
        base_range = {'sheetId': tasks_sheet_id, 'startRowIndex': 1, 'endRowIndex': 2000, 'startColumnIndex': 4, 'endColumnIndex': 5}
        requests.extend([
            {'addConditionalFormatRule': {
                'index': 0,
                'rule': {
                    'ranges': [base_range],
                    'booleanRule': {
                        'condition': {'type': 'TEXT_EQ', 'values': [{'userEnteredValue': 'Open'}]},
                        'format': {'backgroundColor': color(1.0, 0.85, 0.85), 'textFormat': {'bold': True}}
                    }
                }
            }},
            {'addConditionalFormatRule': {
                'index': 0,
                'rule': {
                    'ranges': [base_range],
                    'booleanRule': {
                        'condition': {'type': 'TEXT_EQ', 'values': [{'userEnteredValue': 'Closed'}]},
                        'format': {'backgroundColor': color(0.80, 1.00, 0.80), 'textFormat': {'bold': True}}
                    }
                }
            }},
        ])

    if requests:
        sheets_service.spreadsheets().batchUpdate(spreadsheetId=SPREADSHEET_ID, body={'requests': requests}).execute()

def ensure_tabs_and_headers(sheets_service, project_name):
    tab_tracker = f"{project_name} - Project Tracker"
    tab_tasks = f"{project_name} - Tasks"
    meta = sheets_service.spreadsheets().get(spreadsheetId=SPREADSHEET_ID).execute()
    sheets = meta.get('sheets', [])
    titles = {s['properties']['title']: s['properties']['sheetId'] for s in sheets}

    tracker_created = False
    tasks_created = False

    if tab_tracker not in titles:
        resp = sheets_service.spreadsheets().batchUpdate(
            spreadsheetId=SPREADSHEET_ID, body={'requests':[{'addSheet':{'properties':{'title':tab_tracker}}}]}
        ).execute()
        tracker_sheet_id = resp['replies'][0]['addSheet']['properties']['sheetId']
        titles[tab_tracker] = tracker_sheet_id
        tracker_created = True
    else:
        tracker_sheet_id = titles[tab_tracker]

    if tab_tasks not in titles:
        resp = sheets_service.spreadsheets().batchUpdate(
            spreadsheetId=SPREADSHEET_ID, body={'requests':[{'addSheet':{'properties':{'title':tab_tasks}}}]}
        ).execute()
        titles[tab_tasks] = resp['replies'][0]['addSheet']['properties']['sheetId']
        tasks_created = True

    headers_tracker = ["Project","Source","Type","Part Number","Manufacturer","Quantity","Status","PO Number","Estimated Ship Date (ESD)","Date Received","Date Scheduled","Date Delivered"]
    headers_tasks   = ["Project","Task","Assignee","Due Date","Status","Notes"]

    sheets_service.spreadsheets().values().update(
        spreadsheetId=SPREADSHEET_ID, range=f"{tab_tracker}!A1:L1",
        valueInputOption='RAW', body={'values':[headers_tracker]}
    ).execute()
    sheets_service.spreadsheets().values().update(
        spreadsheetId=SPREADSHEET_ID, range=f"{tab_tasks}!A1:F1",
        valueInputOption='RAW', body={'values':[headers_tasks]}
    ).execute()

    reqs = [
        {'updateSheetProperties': {'properties': {'sheetId': tracker_sheet_id,'gridProperties': {'frozenRowCount':1}}, 'fields':'gridProperties.frozenRowCount'}},
        {'repeatCell': {'range': {'sheetId': tracker_sheet_id,'startRowIndex':1,'endRowIndex':2000,'startColumnIndex':5,'endColumnIndex':6},
                        'cell': {'userEnteredFormat': {'numberFormat': {'type':'NUMBER','pattern':'#,##0'}}},
                        'fields':'userEnteredFormat.numberFormat'}},
        {'updateSheetProperties': {'properties': {'sheetId': titles[tab_tasks],'gridProperties': {'frozenRowCount':1}}, 'fields':'gridProperties.frozenRowCount'}}
    ]
    sheets_service.spreadsheets().batchUpdate(spreadsheetId=SPREADSHEET_ID, body={'requests': reqs}).execute()

    apply_dropdowns(sheets_service, titles, tab_tracker, tab_tasks)
    apply_status_colors(sheets_service, titles, tab_tracker, tab_tasks, tracker_created, tasks_created)
    upsert_dashboard_row(sheets_service, project_name, tracker_sheet_id)

    return tab_tracker

def _is_tax_or_total(line: dict) -> bool:
    for key in ("Type", "Part Number", "Manufacturer"):
        val = _n(line.get(key, ""))
        if any(tok in val for tok in ("tax", "total")):
            return True
    return False

def update_google_sheet(sheets_service, project_name, data_rows):
    tab_tracker = ensure_tabs_and_headers(sheets_service, project_name)
    if not data_rows:
        print("No data to append (tabs were created/updated).")
        return

    filtered = [d for d in data_rows if not _is_tax_or_total(d)]
    if not filtered:
        print("All parsed rows were filtered (tax/total). Nothing to append.")
        return

    values = []
    for d in filtered:
        values.append([
            d.get('Project',''), d.get('Source',''), d.get('Type',''),
            d.get('Part Number',''), d.get('Manufacturer',''), d.get('Quantity',''),
            d.get('Status',''), d.get('PO Number',''), d.get('Estimated Ship Date (ESD)',''),
            d.get('Date Received',''), d.get('Date Scheduled',''), d.get('Date Delivered','')
        ])
    sheets_service.spreadsheets().values().append(
        spreadsheetId=SPREADSHEET_ID, range=f"{tab_tracker}!A2:L",
        valueInputOption='USER_ENTERED', body={'values': values}
    ).execute()
    print(f"Sheet '{tab_tracker}' updated with {len(values)} row(s).")

# ============================ Main =============================
if __name__ == '__main__':
    try:
        import openpyxl  # noqa
    except ImportError:
        print("openpyxl not found. Installing...")
        os.system("pip install openpyxl")
        import openpyxl  # noqa

    gmail_service, sheets_service = setup_services()
    if gmail_service and sheets_service:
        messages = get_emails(gmail_service)
        if messages:
            for m in messages:
                try:
                    project_guess, rows = get_and_process_emails_single(gmail_service, m)
                    if project_guess is None:
                        continue

                    resolution = resolve_okd_project(
                        sheets_service,
                        SPREADSHEET_ID,
                        project_guess,
                    )

                    if not resolution.confirmed:
                        if DEBUG:
                            print(
                                f"[OKD] PROJECT {resolution.status}: {project_guess} "
                                f"method={resolution.method} suggestions={resolution.suggestions}"
                            )
                        # Do not create tabs and leave unread for review.
                        continue

                    project_name = resolution.canonical_project

                    # Ensure every imported row uses the confirmed canonical name.
                    for row in rows:
                        row["Project"] = project_name

                    update_google_sheet(sheets_service, project_name, rows)

                    # Only mark the email read after the confirmed project has been
                    # successfully created/updated.
                    gmail_service.users().messages().modify(
                        userId="me",
                        id=m["id"],
                        body={"removeLabelIds": ["UNREAD"]}
                    ).execute()
                except Exception as e:
                    print(f"Error processing a message; continuing. Details: {e}")
    print("PUMA automation script finished.")
